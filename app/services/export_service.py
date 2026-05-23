from __future__ import annotations

import re
import textwrap
from pathlib import Path

import json
from docx import Document
from docx.shared import Pt


_TIMESTAMP_LINE_RE = re.compile(
    r"^\(\d{1,2}:\d{2}(?::\d{2})?\s+-\s+\d{1,2}:\d{2}(?::\d{2})?\)$"
)


def format_time(seconds: float) -> str:
    total_seconds = max(0, int(seconds))
    h = total_seconds // 3600
    m = (total_seconds % 3600) // 60
    s = total_seconds % 60
    if h > 0:
        return f"{h}:{m:02}:{s:02}"
    return f"{m}:{s:02}"


def format_time_range(start_seconds: float, end_seconds: float | None) -> str:
    end = start_seconds if end_seconds is None else max(start_seconds, end_seconds)
    return f"({format_time(start_seconds)} - {format_time(end)})"


def _timestamped_block(start_seconds: float, end_seconds: float | None, text: str) -> str:
    return f"{format_time_range(start_seconds, end_seconds)}\n{text.strip()}"


def _float_or_default(value, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def get_transcript_content(project, include_timestamps: bool = False) -> str:
    reviewed_transcript = ""
    if project.output and project.output.reviewed_transcript:
        reviewed_transcript = project.output.reviewed_transcript.strip()
    if reviewed_transcript:
        return reviewed_transcript

    if not include_timestamps or not project.chunks:
        if project.output and project.output.full_transcript:
            return project.output.full_transcript
        return ""
        
    lines = []
    for chunk in project.chunks:
        offset = chunk.start_seconds or 0
        fallback_end = chunk.end_seconds if chunk.end_seconds is not None else offset
        
        if not chunk.segments_json:
            if include_timestamps:
                lines.append(
                    _timestamped_block(offset, fallback_end, chunk.transcript_text or "")
                )
            else:
                lines.append(chunk.transcript_text or "")
            continue
            
        try:
            segments = json.loads(chunk.segments_json)
            
            current_paragraph_start = None
            current_paragraph_end = None
            current_paragraph_text = []
            
            for seg in segments:
                segment_start = _float_or_default(seg.get("start"), 0)
                start = offset + segment_start
                end = offset + _float_or_default(seg.get("end"), segment_start)
                text = seg.get("text", "").strip()
                if not text:
                    continue
                    
                if current_paragraph_start is None:
                    current_paragraph_start = start
                    current_paragraph_end = end
                else:
                    current_paragraph_end = max(current_paragraph_end or start, end)
                    
                current_paragraph_text.append(text)
                
                duration = (current_paragraph_end or start) - current_paragraph_start
                if (duration > 30.0 and text[-1] in ".!?。") or duration > 60.0:
                    lines.append(
                        _timestamped_block(
                            current_paragraph_start,
                            current_paragraph_end,
                            " ".join(current_paragraph_text),
                        )
                    )
                    current_paragraph_start = None
                    current_paragraph_end = None
                    current_paragraph_text = []
                    
            if current_paragraph_text:
                lines.append(
                    _timestamped_block(
                        current_paragraph_start,
                        current_paragraph_end,
                        " ".join(current_paragraph_text),
                    )
                )
                
        except Exception:
            lines.append(_timestamped_block(offset, fallback_end, chunk.transcript_text or ""))
            
    return "\n\n".join(lines).strip()


def _strip_markdown_inline(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    return text.strip()


def markdown_to_docx(
    markdown_text: str, output_path: str | Path, promote_first_paragraph: bool = True
) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    title_added = False

    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 4)
            text = _strip_markdown_inline(line[level:].strip())
            if not title_added and level == 1:
                document.add_heading(text, level=0)
                title_added = True
            else:
                document.add_heading(text, level=level)
            continue

        if line.startswith(("- ", "* ")):
            document.add_paragraph(_strip_markdown_inline(line[2:]), style="List Bullet")
            continue

        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        if numbered:
            document.add_paragraph(
                _strip_markdown_inline(numbered.group(1)), style="List Number"
            )
            continue

        if _TIMESTAMP_LINE_RE.match(line):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line)
            run.bold = True
            run.font.size = Pt(9)
            continue

        document.add_paragraph(_strip_markdown_inline(line))

    if not document.paragraphs:
        document.add_paragraph("")
    elif promote_first_paragraph and not title_added:
        document.paragraphs[0].style = document.styles["Title"]

    document.save(output)
    return output


def text_to_pdf(text: str, output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    page_width = 595
    page_height = 842
    margin_x = 54
    margin_y = 54
    font_size = 10.5
    line_height = 14
    max_chars = max(40, int((page_width - (margin_x * 2)) / (font_size * 0.52)))
    lines_per_page = int((page_height - (margin_y * 2)) / line_height)

    logical_lines: list[tuple[str, bool]] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            logical_lines.append(("", False))
            continue

        is_timestamp = bool(_TIMESTAMP_LINE_RE.match(line))
        if is_timestamp:
            logical_lines.append((line, True))
            continue

        wrapped = textwrap.wrap(
            line,
            width=max_chars,
            break_long_words=False,
            replace_whitespace=False,
        )
        for wrapped_line in wrapped or [""]:
            logical_lines.append((wrapped_line, False))

    if not logical_lines:
        logical_lines = [("", False)]

    pages = [
        logical_lines[index : index + lines_per_page]
        for index in range(0, len(logical_lines), lines_per_page)
    ]

    font_regular_id = 3 + len(pages) * 2
    font_bold_id = font_regular_id + 1

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
    ]

    page_ids = [3 + index * 2 for index in range(len(pages))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())

    for index, page_lines in enumerate(pages):
        page_id = 3 + index * 2
        content_id = page_id + 1
        stream_parts = [
            "BT",
            f"{margin_x} {page_height - margin_y} Td",
            f"{line_height} TL",
        ]
        for line, is_timestamp in page_lines:
            if not line:
                stream_parts.append("T*")
                continue
            font_name = "F2" if is_timestamp else "F1"
            size = 9.5 if is_timestamp else font_size
            stream_parts.append(f"/{font_name} {size} Tf")
            stream_parts.append(f"({_pdf_escape(line)}) Tj")
            stream_parts.append("T*")
        stream_parts.append("ET")
        stream = "\n".join(stream_parts).encode("cp1252", errors="replace")

        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {page_width} {page_height}] "
                f"/Resources << /Font << /F1 {font_regular_id} 0 R /F2 {font_bold_id} 0 R >> >> "
                f"/Contents {content_id} 0 R >>"
            ).encode()
        )
        objects.append(
            b"<< /Length "
            + str(len(stream)).encode()
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )

    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")

    output.write_bytes(_build_pdf(objects))
    return output


def markdown_to_pdf(markdown_text: str, output_path: str | Path) -> Path:
    lines: list[str] = []
    for raw_line in markdown_text.splitlines():
        line = raw_line.strip()
        if not line:
            lines.append("")
            continue
        if line.startswith("#"):
            line = _strip_markdown_inline(line.lstrip("#").strip())
        elif line.startswith(("- ", "* ")):
            line = "- " + _strip_markdown_inline(line[2:])
        else:
            numbered = re.match(r"^(\d+\.)\s+(.*)$", line)
            if numbered:
                line = f"{numbered.group(1)} {_strip_markdown_inline(numbered.group(2))}"
            else:
                line = _strip_markdown_inline(line)
        lines.append(line)
    return text_to_pdf("\n".join(lines), output_path)


def _pdf_escape(text: str) -> str:
    text = text.encode("cp1252", errors="replace").decode("cp1252")
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _build_pdf(objects: list[bytes]) -> bytes:
    pdf = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    offsets = [0]

    for object_number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{object_number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_position = len(pdf)
    pdf += f"xref\n0 {len(objects) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        pdf += f"{offset:010} 00000 n \n".encode()

    pdf += (
        b"trailer\n"
        + f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
        + b"startxref\n"
        + str(xref_position).encode()
        + b"\n%%EOF\n"
    )
    return pdf
